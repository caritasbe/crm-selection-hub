#!/usr/bin/env python3
"""
Ingestion de la base documentaire du CRM Selection Hub.

Lit private/manifest.json et les fichiers originaux de private/documents/,
extrait le texte page par page (PDF), diapositive par diapositive (PPTX),
onglet par onglet (XLSX) ou section par section (DOCX), puis produit :

  private/knowledge/documents.json   métadonnées enrichies (nombre de pages, format…)
  private/knowledge/pages/<id>.json  texte par page, pour la visionneuse
  private/knowledge/chunks.json      passages indexés pour la recherche et l'assistant

Prérequis : poppler-utils (pdftotext, pdfinfo), python-docx, python-pptx, openpyxl.
Usage : python3 scripts/ingest.py  (ou npm run ingest)
"""
import json
import os
import re
import subprocess
import sys
import unicodedata
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "private"
DOCS = Path(os.environ.get("DOCS_DIR", DATA / "documents"))
OUT = Path(os.environ.get("KNOWLEDGE_DIR", DATA / "knowledge"))

CHUNK_MAX = 1400
CHUNK_OVERLAP = 220

# Thématiques alignées sur les dimensions du comparateur
THEMES = {
    "marketing-automation": ["journey", "parcours", "workflow", "scénario", "scenario", "automation", "automatis", "trigger", "déclencheur", "nurturing", "campagne automatis"],
    "cdp": ["cdp", "customer data platform", "profil unifié", "unified profile", "data cloud", "identity resolution", "dataverse", "360"],
    "integrations": ["iraiser", "iss", "api", "sftp", "middleware", "n8n", "connecteur", "connector", "webhook", "intégration", "integration", "dynamics", "formidable", "zapier", "make"],
    "architecture": ["architecture", "source de vérité", "source of truth", "base mère", "système de référence", "system of record", "couche", "layer", "hub"],
    "segmentation": ["segment", "rfm", "personnalis", "personalis", "ciblage", "targeting", "scoring", "score", "audience"],
    "consentements": ["consent", "opt-in", "opt in", "préférence", "preference", "rgpd", "gdpr"],
    "migration": ["migration", "import", "dédoublonn", "deduplic", "cadrage", "go-live", "go live", "hypercare", "formation", "training", "planning", "phase", "jalon", "milestone", "jours/homme", "person-days"],
    "support": ["support", "maintenance", "sla", "c-care", "hypercare", "customer success", "csm", "garantie", "uptime"],
    "couts": ["€", "eur", "prix", "price", "coût", "cost", "licence", "license", "budget", "remise", "discount", "tarif", "htva", "crédit", "credit"],
    "contrat": ["contrat", "contract", "engagement", "résili", "terminat", "renouvel", "renewal", "conditions générales", "terms", "préavis", "notice", "commitment", "réversibilité", "reversib", "lock-in", "valid"],
    "securite": ["iso", "sécurité", "security", "hébergement", "hosting", "dpa", "sous-traitant", "sub-processor", "chiffrement", "encryption", "sso", "mfa", "soc 2", "violation", "breach"],
    "collecte": ["quiz", "sondage", "survey", "enquête", "pétition", "petition", "formulaire", "form", "gamif", "landing"],
    "reporting": ["dashboard", "tableau de bord", "reporting", "power bi", "kpi", "data lake", "bigquery", "analytics", "rapport"],
    "ia": [" ia ", " ai ", "agent", "copilot", "breeze", "agentforce", "agentic", "llm", "intelligence artificielle", "artificial intelligence"],
}


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s.lower())
    return "".join(c for c in s if not unicodedata.combining(c))


def clean_lines(text: str) -> list:
    lines = []
    for raw in text.splitlines():
        line = re.sub(r"[ \t ]{2,}", "  ", raw).rstrip()
        line = line.replace("​", "")
        if line.strip():
            lines.append(line.strip())
    return lines


def extract_pdf(path: Path):
    info = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True).stdout
    n = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    pages = []
    for i in range(1, n + 1):
        t = subprocess.run(["pdftotext", "-layout", "-f", str(i), "-l", str(i), str(path), "-"], capture_output=True, text=True).stdout
        pages.append({"page": i, "label": f"p. {i}", "lines": clean_lines(t)})
    return pages


def extract_pptx(path: Path):
    from pptx import Presentation

    pres = Presentation(str(path))
    pages = []
    for i, slide in enumerate(pres.slides, 1):
        buf = []

        def walk(shape):
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    buf.extend(t.splitlines())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    buf.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
            if shape.shape_type == 6:  # groupe
                for sub in shape.shapes:
                    walk(sub)

        for shape in slide.shapes:
            walk(shape)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                buf.append("[Notes] " + notes)
        pages.append({"page": i, "label": f"diapo {i}", "lines": clean_lines("\n".join(buf))})
    return pages


def extract_xlsx(path: Path, exclude):
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    pages = []
    idx = 0
    for ws in wb.worksheets:
        if ws.title in exclude:
            continue
        idx += 1
        lines = []
        for row in ws.iter_rows(values_only=True):
            vals = []
            for v in row:
                if v is None:
                    vals.append("")
                elif isinstance(v, float):
                    vals.append(f"{v:.2f}".rstrip("0").rstrip(".") if abs(v) >= 1 else f"{v:.4g}")
                else:
                    vals.append(str(v).replace("\n", " "))
            while vals and vals[-1] == "":
                vals.pop()
            if any(x.strip() for x in vals):
                lines.append(" | ".join(vals).strip(" |"))
        pages.append({"page": idx, "label": f"onglet « {ws.title} »", "lines": clean_lines("\n".join(lines))})
    return pages


def extract_docx(path: Path):
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    sections = []
    current = {"title": "Page de titre", "lines": []}
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, d)
            text = p.text.strip()
            if not text:
                continue
            style = p.style.name if p.style is not None else ""
            is_heading = (style.startswith("Heading") and style[-1:] in "12") or re.match(r"^Annexe [A-F]\.", text)
            if is_heading:
                if current["lines"]:
                    sections.append(current)
                current = {"title": text, "lines": [text]}
            else:
                current["lines"].append(text)
        elif el.tag == qn("w:tbl"):
            t = Table(el, d)
            for row in t.rows:
                cells = []
                for c in row.cells:
                    tx = c.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != tx:
                        cells.append(tx)
                current["lines"].append(" | ".join(cells))
    if current["lines"]:
        sections.append(current)
    pages = []
    for i, s in enumerate(sections, 1):
        title = re.sub(r"\s+", " ", s["title"])
        pages.append({"page": i, "label": f"§ {title}"[:90], "lines": clean_lines("\n".join(s["lines"]))})
    return pages


def strip_repeated_lines(pages):
    """Supprime les pieds de page et en-têtes répétés (présents sur plus de 60 % des pages)."""
    if len(pages) < 5:
        return pages
    counts = Counter()
    for p in pages:
        for l in set(p["lines"]):
            counts[l] += 1
    threshold = max(3, int(len(pages) * 0.6))
    repeated = {l for l, c in counts.items() if c >= threshold and len(l) > 12}
    for p in pages:
        p["lines"] = [l for l in p["lines"] if l not in repeated]
    return pages


def detect_themes(text: str):
    t = " " + norm(text) + " "
    found = []
    for theme, kws in THEMES.items():
        score = sum(t.count(norm(k)) for k in kws)
        if score >= 2 or (score >= 1 and theme in ("collecte", "ia", "cdp")):
            found.append(theme)
    return found


def chunk_page(lines):
    chunks, buf, size = [], [], 0
    for line in lines:
        if size + len(line) > CHUNK_MAX and buf:
            chunks.append("\n".join(buf))
            # chevauchement : on reprend la fin du passage précédent
            overlap, osize = [], 0
            for l in reversed(buf):
                if osize + len(l) > CHUNK_OVERLAP:
                    break
                overlap.insert(0, l)
                osize += len(l)
            buf, size = overlap[:], osize
        buf.append(line)
        size += len(line) + 1
    if buf:
        chunks.append("\n".join(buf))
    return chunks


def main():
    manifest = json.loads((DATA / "manifest.json").read_text(encoding="utf-8"))
    (OUT / "pages").mkdir(parents=True, exist_ok=True)
    documents, all_chunks = [], []
    for meta in manifest["documents"]:
        path = DOCS / meta["file"]
        if not path.exists():
            print(f"!! fichier manquant : {path}", file=sys.stderr)
            continue
        ext = path.suffix.lower()
        if ext == ".pdf":
            pages = extract_pdf(path)
        elif ext == ".pptx":
            pages = extract_pptx(path)
        elif ext == ".xlsx":
            pages = extract_xlsx(path, set(meta.get("excludeSheets", [])))
        elif ext == ".docx":
            pages = extract_docx(path)
        else:
            print(f"!! format non pris en charge : {path}", file=sys.stderr)
            continue
        pages = strip_repeated_lines(pages)
        page_out = [{"page": p["page"], "label": p["label"], "text": "\n".join(p["lines"])} for p in pages]
        (OUT / "pages" / f"{meta['id']}.json").write_text(json.dumps(page_out, ensure_ascii=False), encoding="utf-8")

        doc_themes = Counter()
        n_chunks = 0
        seen_bodies = set()
        no_index = set(meta.get("noIndexPages", []))
        for p in pages:
            if not p["lines"] or p["label"] in no_index:
                continue
            for j, text in enumerate(chunk_page(p["lines"])):
                if len(text.strip()) < 25:
                    continue
                # pages au contenu identique (ex. feuilles de notation vierges) : un seul exemplaire indexé
                body = norm("\n".join(text.splitlines()[1:]) if text.count("\n") > 2 else text)
                if body in seen_bodies:
                    continue
                seen_bodies.add(body)
                themes = detect_themes(text)
                doc_themes.update(themes)
                all_chunks.append({
                    "id": f"{meta['id']}#{p['page']}.{j}",
                    "docId": meta["id"],
                    "page": p["page"],
                    "label": p["label"],
                    "vendors": meta["vendors"],
                    "themes": themes,
                    "text": text,
                })
                n_chunks += 1
        doc = dict(meta)
        doc.update({
            "format": ext[1:],
            "pageCount": len(pages),
            "locatorType": {"pdf": "page", "pptx": "diapositive", "xlsx": "onglet", "docx": "section"}[ext[1:]],
            "pages": [{"page": p["page"], "label": p["label"]} for p in pages],
            "sizeBytes": path.stat().st_size,
            "chunkCount": n_chunks,
            "themes": [t for t, c in doc_themes.most_common() if c >= 2][:8],
        })
        documents.append(doc)
        print(f"{meta['id']:<16} {ext[1:]:<5} {len(pages):>3} pages  {n_chunks:>4} passages")

    (OUT / "documents.json").write_text(json.dumps(documents, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUT / "chunks.json").write_text(json.dumps(all_chunks, ensure_ascii=False), encoding="utf-8")
    print(f"\n{len(documents)} documents, {len(all_chunks)} passages -> {OUT}")


if __name__ == "__main__":
    main()
