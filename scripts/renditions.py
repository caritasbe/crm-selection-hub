#!/usr/bin/env python3
"""
Prépare l'affichage visuel des documents (dossier private/renditions, exclu de Git) :
  <id>.pdf  version PDF des fichiers Office (PowerPoint, Word, Excel) via LibreOffice ;
  <id>.jpg  vignette de la première page, pour la bibliothèque.
Les onglets exclus (excludeSheets) ne figurent pas dans la version PDF des classeurs.
Un document inchangé n'est pas reconverti.

  python3 scripts/renditions.py        (lancé aussi par npm run ingest)
"""
import hashlib, json, os, shutil, subprocess, sys, tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PRIV = ROOT / "private"
DOCS = PRIV / "documents"
OUT = PRIV / "renditions"


def digest(path: Path, extra: str = "") -> str:
    h = hashlib.sha256(extra.encode())
    h.update(path.read_bytes())
    return h.hexdigest()[:24]


def prepare_xlsx(src: Path, exclude, tmp: Path) -> Path:
    from openpyxl import load_workbook
    wb = load_workbook(src)
    for name in list(wb.sheetnames):
        if name in exclude:
            del wb[name]
            continue
        ws = wb[name]
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
    dst = tmp / src.name
    wb.save(dst)
    return dst


def to_pdf(src: Path, tmp: Path) -> Path:
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    pdf = tmp / (src.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError(f"conversion échouée : {src.name}")
    return pdf


def thumb(pdf: Path, dst: Path):
    base = dst.with_suffix("")
    subprocess.run(["pdftoppm", "-f", "1", "-l", "1", "-singlefile", "-scale-to", "640", "-jpeg", "-jpegopt", "quality=78", str(pdf), str(base)],
                   check=True, timeout=120)


def main():
    manifest = json.loads((PRIV / "manifest.json").read_text(encoding="utf-8"))
    OUT.mkdir(parents=True, exist_ok=True)
    stamps_path = OUT / "stamps.json"
    stamps = json.loads(stamps_path.read_text()) if stamps_path.exists() else {}
    done = 0
    for meta in manifest["documents"]:
        src = DOCS / meta["file"]
        if not src.exists():
            print(f"  absent : {meta['file']}")
            continue
        did = meta["id"]
        stamp = digest(src, json.dumps(meta.get("excludeSheets", [])))
        is_pdf = src.suffix.lower() == ".pdf"
        pdf_out = OUT / f"{did}.pdf"
        jpg_out = OUT / f"{did}.jpg"
        if stamps.get(did) == stamp and jpg_out.exists() and (is_pdf or pdf_out.exists()):
            continue
        with tempfile.TemporaryDirectory() as t:
            tmp = Path(t)
            if is_pdf:
                pdf = src
                if pdf_out.exists():
                    pdf_out.unlink()
            else:
                s = prepare_xlsx(src, set(meta.get("excludeSheets", [])), tmp) if src.suffix.lower() == ".xlsx" else src
                pdf = to_pdf(s, tmp)
                shutil.copyfile(pdf, pdf_out)
            thumb(pdf, jpg_out)
        stamps[did] = stamp
        done += 1
        print(f"  {did} : {'vignette' if is_pdf else 'PDF + vignette'}")
    stamps_path.write_text(json.dumps(stamps, indent=1))
    print(f"Rendus : {done} document(s) traité(s), {len(manifest['documents']) - done} inchangé(s).")


if __name__ == "__main__":
    main()
